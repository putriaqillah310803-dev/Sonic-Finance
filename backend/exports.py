"""Excel and PDF export helpers for SonicGo reports."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def _rp(v):
    try:
        return "Rp {:,.0f}".format(float(v)).replace(",", ".")
    except Exception:
        return str(v)


def build_excel(title, meta, sections):
    """sections: list of dict {name, headers:[], rows:[[...]]}"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    red = PatternFill("solid", fgColor="DC2626")
    light = PatternFill("solid", fgColor="FEF2F2")
    white_bold = Font(bold=True, color="FFFFFF", size=12)
    bold = Font(bold=True)
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(row=r, column=1, value=title)
    c.font = Font(bold=True, size=16, color="DC2626")
    r += 1
    for k, v in meta.items():
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1

    for sec in sections:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, len(sec["headers"])))
        hc = ws.cell(row=r, column=1, value=sec["name"])
        hc.fill = red
        hc.font = white_bold
        r += 1
        for i, h in enumerate(sec["headers"], start=1):
            cell = ws.cell(row=r, column=i, value=h)
            cell.font = bold
            cell.fill = light
            cell.border = border
        r += 1
        for row in sec["rows"]:
            for i, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=i, value=val)
                cell.border = border
            r += 1
        r += 1

    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_pdf(title, meta, sections):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], textColor=colors.HexColor("#DC2626"), fontSize=18)
    elems = [Paragraph(title, title_style), Spacer(1, 6)]
    meta_txt = "  |  ".join(f"<b>{k}:</b> {v}" for k, v in meta.items())
    elems.append(Paragraph(meta_txt, styles["Normal"]))
    elems.append(Spacer(1, 14))

    for sec in sections:
        elems.append(Paragraph(f"<b>{sec['name']}</b>", styles["Heading3"]))
        data = [sec["headers"]] + sec["rows"]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DC2626")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FEF2F2")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elems.append(t)
        elems.append(Spacer(1, 16))

    doc.build(elems)
    buf.seek(0)
    return buf
